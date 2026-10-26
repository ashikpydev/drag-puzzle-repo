"""Build the sample form and its media by driving the real puzzle maker in a headless browser.
This doubles as an end-to-end test of the maker's zip export.
Usage: python3 tools/make_sample_form.py
"""
import os, shutil, zipfile
from openpyxl import Workbook
from playwright.sync_api import sync_playwright

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, '..')
MAKER = 'file://' + os.path.abspath(os.path.join(ROOT, 'extras', 'puzzle-maker', 'puzzle-maker.html'))
MEDIA = os.path.join(ROOT, 'extras', 'sample-form', 'media')
TMP = os.path.join(HERE, '_tmp')
shutil.rmtree(TMP, ignore_errors=True); os.makedirs(TMP)
shutil.rmtree(MEDIA, ignore_errors=True); os.makedirs(MEDIA)

# centre positions in percent of the picture
ITEMS = [
    dict(pic='ball', name='puzzle_practice', practice=True, theme='sunny', keep_wrong=2,
         gap=(50, 46.7), wrong=[(50, 85), (17.7, 17.7)]),
    dict(pic='fish', name='puzzle_1', practice=False, theme='sky', keep_wrong=3,
         gap=(54.3, 42.7), wrong=[(18.3, 42.7), (21, 85), (41.7, 44.3)]),
    dict(pic='house', name='puzzle_2', practice=False, theme='forest', keep_wrong=3,
         gap=(53.7, 61), wrong=[(50, 34), (39.7, 56), (87.7, 54.3)]),
]
SEEDS = [(15, 15), (81, 15), (15, 83)]   # centres of the three automatically placed wrong boxes


def drag(pg, cv, frm, to):
    b = cv.bounding_box()
    fx, fy = b['x'] + frm[0] / 100 * b['width'], b['y'] + frm[1] / 100 * b['height']
    tx, ty = b['x'] + to[0] / 100 * b['width'], b['y'] + to[1] / 100 * b['height']
    pg.mouse.move(fx, fy); pg.mouse.down()
    for i in range(1, 11):
        pg.mouse.move(fx + (tx - fx) * i / 10, fy + (ty - fy) * i / 10)
    pg.mouse.up()


survey_rows, choice_rows = [], []
with sync_playwright() as p:
    br = p.chromium.launch()
    ctx = br.new_context(accept_downloads=True, viewport={'width': 1300, 'height': 1000})
    for it in ITEMS:
        pg = ctx.new_page()
        pg.goto(MAKER)
        pg.set_input_files('#file', os.path.join(HERE, 'pictures', it['pic'] + '.png'))
        pg.wait_for_timeout(400)
        pg.fill('#name', it['name'])
        pg.select_option('#theme', it['theme'])
        if it['practice']:
            pg.check('#prac')
        while len(pg.query_selector_all('#rem')) and it['keep_wrong'] < 3:
            pg.click('#rem'); it['keep_wrong'] += 1
        cv = pg.query_selector('#cv')
        drag(pg, cv, (51, 46), it['gap'])
        for seed, target in zip(SEEDS, it['wrong']):
            drag(pg, cv, seed, target)
        pg.wait_for_timeout(200)
        rows = pg.input_value('#rows').split('\n')
        with pg.expect_download() as dl:
            pg.click('#zip')
        path = os.path.join(TMP, it['name'] + '.zip')
        dl.value.save_as(path)
        with zipfile.ZipFile(path) as z:
            assert z.testzip() is None
            for n in z.namelist():
                if n.endswith('.png'):
                    z.extract(n, MEDIA)
        # parse the rows text
        s_i = rows.index('SURVEY SHEET (one row)')
        c_i = rows.index('CHOICES SHEET (rows)')
        survey_rows.append(rows[s_i + 2].split('\t'))
        survey_hdr = rows[s_i + 1].split('\t')
        choice_hdr = rows[c_i + 1].split('\t')
        for r in rows[c_i + 2:]:
            if r.strip():
                choice_rows.append(r.split('\t'))
        pg.close()
    br.close()

wb = Workbook()
ws = wb.active; ws.title = 'survey'
ws.append(survey_hdr)
for r in survey_rows:
    ws.append(r)
# optional log fields: copy the plug-in log into the data (one simple expression each)
for r in survey_rows[1:]:
    ws.append(['calculate', r[1] + '_log', '', '', '', '', 'plug-in-metadata(${%s})' % r[1]])
ws.cell(row=1, column=7, value='calculation')
ch = wb.create_sheet('choices')
ch.append(choice_hdr)
for r in choice_rows:
    ch.append(r)
st = wb.create_sheet('settings')
st.append(['form_title', 'form_id', 'version'])
st.append(['Drag puzzle sample form', 'drag_puzzle_sample', 2026102701])
for w in (ws, ch, st):
    for col in w.columns:
        w.column_dimensions[col[0].column_letter].width = 38
out = os.path.join(ROOT, 'extras', 'sample-form', 'drag-puzzle-sample-form.xlsx')
wb.save(out)
shutil.rmtree(TMP, ignore_errors=True)
print('saved', out)
print(sorted(os.listdir(MEDIA)))
for r in survey_rows:
    print(r[0], '|', r[1], '|', r[3], '|', r[4])
