"""Test harness: runs the plug-in in a browser with a mock of the SurveyCTO plug-in API.

It reads the real sample form (extras/sample-form/drag-puzzle-sample-form.xlsx) so the tests use exactly the
values SurveyCTO would pass to the plug-in: the question image, the choices with their images, and the
parameters written in the appearance column.
"""
import json
import os
import re
import shutil

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, '..')
SRC = os.path.join(ROOT, 'source')
FORM = os.path.join(ROOT, 'extras', 'sample-form', 'drag-puzzle-sample-form.xlsx')
MEDIA = os.path.join(ROOT, 'extras', 'sample-form', 'media')
OUT = os.path.join(HERE, '_out')

TEMPLATE = open(os.path.join(SRC, 'template.html')).read()
CSS = open(os.path.join(SRC, 'style.css')).read()
SCRIPT = open(os.path.join(SRC, 'script.js')).read()


def reset_out():
    shutil.rmtree(OUT, ignore_errors=True)
    os.makedirs(OUT)
    for f in os.listdir(MEDIA):
        shutil.copy(os.path.join(MEDIA, f), OUT)


def parse_params(appearance):
    m = re.match(r'\s*custom-drag-puzzle\((.*)\)\s*$', appearance or '', re.S)
    params = {}
    if m:
        for part in m.group(1).split(','):
            if '=' in part:
                k, v = part.split('=', 1)
                params[k.strip()] = v.strip().strip("'")
    return params


def sample_items():
    wb = load_workbook(FORM)
    survey = wb['survey']
    hdr = [c.value for c in survey[1]]
    choices_ws = wb['choices']
    ch_hdr = [c.value for c in choices_ws[1]]
    lists = {}
    for row in choices_ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(ch_hdr, row))
        if d.get('list_name'):
            lists.setdefault(d['list_name'], []).append(
                {'value': str(d['name']), 'label': d.get('label') or '', 'image': d.get('image') or ''})
    items = {}
    for row in survey.iter_rows(min_row=2, values_only=True):
        d = dict(zip(hdr, row))
        t = d.get('type') or ''
        if t.startswith('select_one '):
            items[d['name']] = {
                'name': d['name'],
                'label': d.get('label') or '',
                'scene': d.get('media:image') or '',
                'params': parse_params(d.get('appearance')),
                'choices': lists[t.split(' ', 1)[1]],
            }
    return items


def page_url(key, item, params=None, current='', readonly=False, audio='', scene=None, choices=None, meta=''):
    """Write a harness page and return its file:// URL. params overrides/adds parameters (None removes one)."""
    p = dict(item['params'])
    for k, v in (params or {}).items():
        if v is None:
            p.pop(k, None)
        else:
            p[k] = v
    ch = choices if choices is not None else item['choices']
    fp = {
        'FIELDTYPE': 'select_one', 'LABEL': item['label'], 'HINT': '', 'READONLY': readonly,
        'CURRENT_ANSWER': current, 'MEDIAIMAGE': item['scene'] if scene is None else scene,
        'MEDIAAUDIO': audio,
        'CHOICES': [{'CHOICE_VALUE': c['value'], 'CHOICE_INDEX': i, 'CHOICE_LABEL': c['label'],
                     'CHOICE_SELECTED': False, 'CHOICE_IMAGE': c['image']} for i, c in enumerate(ch)],
    }
    mock = """
    window.__answer = %s; window.__meta = %s; window.__params = %s;
    window.fieldProperties = %s;
    window.getPluginParameter = function(n){ return window.__params[n]; };
    window.setAnswer = function(v){ window.__answer = v; window.fieldProperties.CURRENT_ANSWER = v; };
    window.setMetaData = function(v){ window.__meta = v; };
    window.getMetaData = function(){ return window.__meta; };
    """ % (json.dumps(current), json.dumps(meta), json.dumps(p), json.dumps(fp))
    tpl = TEMPLATE.replace('{{{LABEL}}}', item['label']).replace('{{{HINT}}}', '')
    html = ('<!doctype html><html><head><meta name="viewport" content="width=device-width, initial-scale=1">'
            '<style>body{margin:0;padding:12px;font-family:sans-serif}' + CSS + '</style></head><body>'
            + tpl + '<script>' + mock + '</script><script>' + SCRIPT + '</script></body></html>')
    path = os.path.join(OUT, key + '.html')
    open(path, 'w').write(html)
    return 'file://' + path


def center(el):
    b = el.bounding_box()
    return b['x'] + b['width'] / 2, b['y'] + b['height'] / 2


def gap_center(pg, params):
    """Centre of the gap in screen pixels (pictures in the sample form are square)."""
    sc = pg.query_selector('#dp-scene').bounding_box()
    gx, gy, gw = float(params['gap_x']), float(params['gap_y']), float(params.get('gap_w', 22))
    gh = float(params.get('gap_h', gw))
    return (sc['x'] + (gx + gw / 2) / 100 * sc['width'], sc['y'] + (gy + gh / 2) / 100 * sc['height'])


def drag_to(pg, el, tx, ty, wait=500):
    x, y = center(el)
    pg.mouse.move(x, y)
    pg.mouse.down()
    for i in range(1, 13):
        pg.mouse.move(x + (tx - x) * i / 12, y + (ty - y) * i / 12)
    pg.mouse.up()
    pg.wait_for_timeout(wait)


def piece(pg, value):
    return pg.query_selector('.dp-piece[data-id="%s"]' % value)
